from flask import Flask, render_template, request, jsonify, send_file
import ollama
import json
import csv
import io
import os
import re
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Google API Client Imports
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build

app = Flask(__name__)

OLLAMA_MODEL = "qwen2.5vl:7b"

# Scopes required to create and format Google Sheets
SCOPES = ['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive.file']

INSTRUCTIONS = """Extract EVERY row from this timesheet. Return ONLY JSON:
{
  "employee_name": "Full Name",
  "rows": [
    {"Date": "Feb 2, 2026", "Time_In": "", "Time_Out": "", "Remarks": "SUNDAY"}
  ]
}
Rules: Match each date to times on SAME line. If no times, leave empty. Clean name matches to normal capitalization."""

ALLOWED_EXTENSIONS = {'jpg', 'jpeg', 'png', 'gif', 'webp'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_attendance_from_image(image_path):
    """Use Ollama with an expanded context window to extract attendance data"""
    try:
        print(f"Processing: {image_path}")
        print(f"Calling Ollama ({OLLAMA_MODEL})...")
        
        with open(image_path, 'rb') as f:
            import base64
            img_b64 = base64.b64encode(f.read()).decode()
        
        response = ollama.generate(
            model=OLLAMA_MODEL,
            prompt=INSTRUCTIONS,
            images=[img_b64],
            stream=False,
            options={
                "num_ctx": 8192  # Safety cushion buffer
            }
        )
        
        raw = response['response'].strip()
        
        json_match = re.search(r'\{.*\}', raw, re.DOTALL)
        if not json_match:
            return {"success": False, "error": "No JSON found in response"}
        
        json_str = json_match.group(0)
        json_str = re.sub(r',\s*}', '}', json_str)
        json_str = re.sub(r',\s*]', ']', json_str)
        
        data = json.loads(json_str)
        return {
            "success": True, 
            "rows": data.get("rows", []), 
            "employee_name": data.get("employee_name", "").strip()
        }
        
    except Exception as e:
        print(f"Ollama Vision Extraction Error: {e}")
        return {"success": False, "error": str(e)}

def build_continuous_rows(table_rows):
    """Fill gaps between dates cleanly across calendar days"""
    row_dict = {}
    valid_dates = []
    
    for row in table_rows:
        date_str = row.get("Date", "").strip()
        dt = None
        for fmt in ("%B %d, %Y", "%b %d, %Y", "%B %d", "%m/%d/%Y", "%d/%m/%Y"):
            try:
                if "2026" not in date_str:
                    dt = datetime.strptime(f"{date_str}, 2026", "%B %d, %Y")
                else:
                    dt = datetime.strptime(date_str, fmt)
                break
            except:
                continue
        
        if dt:
            row_dict[dt.date()] = row
            valid_dates.append(dt.date())
    
    continuous = []
    if valid_dates:
        current = min(valid_dates)
        end = max(valid_dates)
        while current <= end:
            if current in row_dict:
                continuous.append((current, row_dict[current]))
            else:
                continuous.append((current, {
                    "Date": current.strftime("%B %d, %Y"),
                    "Time_In": "", "Time_Out": "", "Remarks": ""
                }))
            current += timedelta(days=1)
    else:
        continuous = [(None, r) for r in table_rows]
    
    return continuous

def normalize_entries(continuous_rows):
    """Format entries with weekend handling"""
    entries = []
    for dt, row in continuous_rows:
        date_val = row.get("Date", "")
        t_in = (row.get("Time_In") or "").strip()
        t_out = (row.get("Time_Out") or "").strip()
        remarks = (row.get("Remarks") or "").strip()
        
        day_name = dt.strftime("%A").upper() if dt else ""
        
        if not t_in and not t_out and day_name in ["SATURDAY", "SUNDAY"]:
            remarks = day_name
        
        is_non_working = (not t_in and not t_out and
                          any(k in remarks.upper() for k in ["SUNDAY", "SATURDAY", "HOLIDAY"]))
        
        if not remarks and not is_non_working and (t_in or t_out):
            remarks = "Regular Working Days"
        
        entries.append({
            "date": date_val,
            "time_in": t_in if not is_non_working else "",
            "time_out": t_out,
            "absent": "" if is_non_working else "0:00",
            "late": "" if is_non_working else "0:00",
            "remarks": remarks,
            "non_working": is_non_working
        })
    return entries

# --- FLASK ROUTES CONFIGURATION ---

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/process', methods=['POST'])
def process_images():
    try:
        if 'files' not in request.files:
            return jsonify({"success": False, "error": "No files selected"}), 400
        
        files = request.files.getlist('files')
        all_entries = []
        detected_employee_name = ''
        
        for file in files:
            if not file.filename or not allowed_file(file.filename):
                continue
            
            temp_path = f"/tmp/{file.filename}"
            os.makedirs(os.path.dirname(temp_path), exist_ok=True)
            file.save(temp_path)
            
            result = extract_attendance_from_image(temp_path)
            
            if result["success"]:
                if result.get("employee_name"):
                    detected_employee_name = result["employee_name"]
                continuous = build_continuous_rows(result["rows"])
                entries = normalize_entries(continuous)
                all_entries.extend(entries)
            
            try:
                os.remove(temp_path)
            except:
                pass
        
        if not all_entries:
            return jsonify({"success": False, "error": "No data could be processed"}), 500
        
        return jsonify({
            "success": True,
            "detected_name": detected_employee_name,
            "total_entries": len(all_entries),
            "entries": all_entries
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/export-google-sheets', methods=['POST'])
def export_google_sheets():
    """Generates an institutional spreadsheet directly in Google Sheets"""
    try:
        data = request.json
        entries = data.get('entries', [])
        employee_name = data.get('employee_name', 'Unknown Employee')
        designation = data.get('designation', 'Project Admin Officer I')
        signee_name = data.get('signee_name', '')
        signee_title = data.get('signee_title', '')

        date_range_str = "February 16 - March 15, 2026"
        if entries:
            date_range_str = f"{entries[0].get('date', '')} - {entries[-1].get('date', '')}"

        if not os.path.exists('credentials.json'):
            return jsonify({"success": False, "error": "Google credentials.json file missing from workspace folder"}), 500
            
        flow = InstalledAppFlow.from_client_secrets_file('credentials.json', SCOPES)
        creds = flow.run_local_server(port=0)
        service = build('sheets', 'v4', credentials=creds)

        spreadsheet_body = {
            'properties': {'title': f"{employee_name}_DTR_{datetime.now().strftime('%Y%m%d')}"}
        }
        ss = service.spreadsheets().create(body=spreadsheet_body, fields='spreadsheetId').execute()
        ss_id = ss.get('spreadsheetId')

        values = [
            ["", "ATENEO DE MANILA UNIVERSITY"],
            ["", "EMPLOYEE's DAILY TIME RECORD"],
            [],
            ["Month:", date_range_str],
            ["Name:", employee_name],
            ["Designation:", designation],
            [],
            ['Date', 'Time In', 'Time Out', 'Absent', 'Late', 'Remarks']
        ]

        row_mapping = {}
        idx = 9
        for e in entries:
            if e.get('non_working'):
                values.append([e.get('date', ''), e.get('remarks', '').upper(), "", "", "", ""])
                row_mapping[idx] = "non_working"
            else:
                values.append([
                    e.get('date', ''), e.get('time_in', ''), e.get('time_out', ''),
                    e.get('absent', '0:00') or "0:00", e.get('late', '0:00') or "0:00", e.get('remarks', '')
                ])
                row_mapping[idx] = "working"
            idx += 1

        values.extend([
            [], [],
            ["Certified true and correct:"],
            ["", "", signee_name],
            ["", "", signee_title]
        ])

        body = {'values': values}
        service.spreadsheets().values().update(
            spreadsheetId=ss_id, range="Sheet1!A1",
            valueInputOption="USER_ENTERED", body=body).execute()

        requests = [
            {
                "updateDimensionProperties": {
                    "range": {"sheetId": 0, "dimension": "COLUMNS", "startIndex": 0, "endIndex": 6},
                    "properties": {"pixelSize": 130},
                    "fields": "pixelSize"
                }
            },
            {"mergeCells": {"range": {"sheetId": 0, "startRowIndex": 1, "endRowIndex": 2, "startColumnIndex": 1, "endColumnIndex": 5}, "mergeType": "MERGE_ALL"}},
            {"mergeCells": {"range": {"sheetId": 0, "startRowIndex": 2, "endRowIndex": 3, "startColumnIndex": 1, "endColumnIndex": 5}, "mergeType": "MERGE_ALL"}},
            {
                "repeatCell": {
                    "range": {"sheetId": 0, "startRowIndex": 1, "endRowIndex": 3, "startColumnIndex": 1, "endColumnIndex": 5},
                    "cell": {"userEnteredFormat": {"textFormat": {"fontFamily": "Times New Roman", "fontSize": 12, "bold": True}}},
                    "fields": "userEnteredFormat.textFormat"
                }
            },
            {
                "repeatCell": {
                    "range": {"sheetId": 0, "startRowIndex": 7, "endRowIndex": 8, "startColumnIndex": 0, "endColumnIndex": 6},
                    "cell": {
                        "userEnteredFormat": {
                            "backgroundColor": {"red": 0.788, "green": 0.854, "blue": 0.972},
                            "textFormat": {"fontFamily": "Arial", "fontSize": 10, "bold": True},
                            "horizontalAlignment": "CENTER"
                        }
                    },
                    "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment)"
                }
            }
        ]

        for r_idx, r_type in row_mapping.items():
            if r_type == "non_working":
                requests.append({
                    "mergeCells": {
                        "range": {"sheetId": 0, "startRowIndex": r_idx - 1, "endRowIndex": r_idx, "startColumnIndex": 1, "endColumnIndex": 6},
                        "mergeType": "MERGE_ALL"
                    }
                })
                requests.append({
                    "repeatCell": {
                        "range": {"sheetId": 0, "startRowIndex": r_idx - 1, "endRowIndex": r_idx, "startColumnIndex": 1, "endColumnIndex": 6},
                        "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER", "textFormat": {"bold": True}}},
                        "fields": "userEnteredFormat(horizontalAlignment,textFormat)"
                    }
                })

        service.spreadsheets().batchUpdate(spreadsheetId=ss_id, body={'requests': requests}).execute()
        return jsonify({"success": True, "sheet_url": f"https://docs.google.com/spreadsheets/d/{ss_id}"})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/export-excel', methods=['POST'])
def export_excel():
    """Outputs an institutional spreadsheet adhering precisely to user dropdown entries"""
    try:
        data = request.json
        entries = data.get('entries', [])
        
        employee_name = data.get('employee_name', 'Unknown Employee')
        designation = data.get('designation', 'Unknown Designation')
        signee_name = data.get('signee_name', 'Signee')
        signee_title = data.get('signee_title', 'Signee Title')
        
        date_range_str = "February 16 - March 15, 2026"
        if entries:
            try:
                date_range_str = f"{entries[0].get('date', '')} - {entries[-1].get('date', '')}"
            except:
                pass

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DTR"
        
        ws.views.sheetView[0].showGridLines = True

        font_title = Font(name="Times New Roman", size=12, bold=True)
        font_regular_bold = Font(name="Arial", size=10, bold=True)
        font_data = Font(name="Arial", size=10, bold=False)
        font_italic = Font(name="Arial", size=10, italic=True)
        fill_header = PatternFill(start_color="C9DAF8", end_color="C9DAF8", fill_type="solid")
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        border_thin = Side(border_style="thin", color="000000")
        box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        ws['B2'] = "ATENEO DE MANILA UNIVERSITY"
        ws['B2'].font = font_title
        ws.merge_cells('B2:E2')
        
        ws['B3'] = "EMPLOYEE's DAILY TIME RECORD"
        ws['B3'].font = font_title
        ws.merge_cells('B3:E3')

        meta_items = [
            ("A5", "B5", "Month:", date_range_str),
            ("A6", "B6", "Name:", employee_name),
            ("A7", "B7", "Designation:", designation)
        ]
        
        for label_cell, val_cell, label, val in meta_items:
            ws[label_cell] = label
            ws[label_cell].font = font_regular_bold
            ws[val_cell] = val
            ws[val_cell].font = font_data

        headers = ['Date', 'Time In', 'Time Out', 'Absent', 'Late', 'Remarks']
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=9, column=col_idx, value=header)
            cell.font = font_regular_bold
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = box_border

        current_row = 10
        for e in entries:
            date_val = e.get('date', '')
            date_cell = ws.cell(row=current_row, column=1, value=date_val)
            date_cell.font = font_data
            date_cell.alignment = align_center
            date_cell.border = box_border
            
            if e.get('non_working'):
                label_text = e.get('remarks', 'NON-WORKING').upper()
                merged_label_cell = ws.cell(row=current_row, column=2, value=label_text)
                merged_label_cell.font = font_regular_bold
                merged_label_cell.alignment = align_center
                
                ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=6)
                for c in range(2, 7):
                    ws.cell(row=current_row, column=c).border = box_border
            else:
                row_data = [
                    e.get('time_in', ''),
                    e.get('time_out', ''),
                    e.get('absent', '0:00') if e.get('absent') != "" else "0:00",
                    e.get('late', '0:00') if e.get('late') != "" else "0:00",
                    e.get('remarks', 'Regular Working Days')
                ]
                for offset, val in enumerate(row_data, 2):
                    cell = ws.cell(row=current_row, column=offset, value=val)
                    cell.font = font_data
                    cell.alignment = align_center if offset < 6 else align_left
                    cell.border = box_border
            current_row += 1

        sig_start = current_row + 2
        ws.cell(row=sig_start, column=1, value="Certified true and correct:").font = font_italic
        
        ws.cell(row=sig_start, column=3, value=employee_name).font = font_regular_bold
        ws.cell(row=sig_start+1, column=3, value=designation).font = font_data
        
        ws.cell(row=sig_start+4, column=3, value=signee_name).font = font_regular_bold
        ws.cell(row=sig_start+5, column=3, value=signee_title).font = font_data

        column_widths = {'A': 16, 'B': 13, 'C': 13, 'D': 12, 'E': 12, 'F': 24}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        excel_file = io.BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        
        filename = f"{employee_name.replace(' ', '_')}_DTR_{datetime.now().strftime('%Y%m%d')}.xlsx"
        return send_file(excel_file, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=filename)
    except Exception as e:
        print(f"Excel Export Engine Error: {e}")
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/health', methods=['GET'])
def health_check():
    """Sync smoothly with JavaScript tracking properties and output errors to terminal"""
    try:
        # Check if Ollama service is reachable
        models_list = ollama.list()
        available_models = [m.get('model', '') for m in models_list.get('models', [])]
        model_ready = any(OLLAMA_MODEL in name for name in available_models) or len(available_models) > 0
        
        # Verify if credentials file is visible relative to execution path
        creds_exist = os.path.exists('credentials.json')
        print(f"--- [HEALTH CHECK] Credentials file found: {creds_exist} ---")
        
        return jsonify({
            "ollama_running": True, 
            "model_ready": model_ready,
            "model": OLLAMA_MODEL,
            "credentials_valid": creds_exist
        })
    except Exception as e:
        # This will print the precise crash reason to your terminal window!
        print(f"--- [HEALTH CHECK CRASH]: {str(e)} ---")
        return jsonify({
            "ollama_running": False, 
            "model_ready": False, 
            "model": OLLAMA_MODEL,
            "error": str(e)
        }), 503

# Keep this statement at the absolute bottom of your script file execution tree
if __name__ == '__main__':
    app.run(debug=True, host='127.0.0.1', port=5000)