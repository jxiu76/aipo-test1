import json
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime, timedelta

# Import the Google GenAI SDK
from google import genai
from google.genai import types

API_KEY = "YOUR_API_KEY"

print("Reading attendance image and preparing to send...")

try:
    with open("attendance.jpg", "rb") as image_file:
        image_bytes = image_file.read()
except FileNotFoundError:
    print("Error: Could not find 'attendance.jpg'. Please make sure it is in the same folder.")
    sys.exit(1)

instructions = """
Analyze this image as a handwritten employee timesheet. 
Extract EVERY single row strictly line-by-line into an array of objects assigned to a top-level key named "rows".

Follow these strict rules to prevent row-shifting errors:
1. Examine each row as an isolated horizontal line. Match the Date directly to the Time In and Time Out written on that exact same line.
2. If a specific Date line has NO handwritten times written next to it (even if it has text in the Remarks column), you MUST set "Time_In": "" and "Time_Out": "". Do NOT pull or shift timestamps from the lines below or above it.
3. Extract the Date exactly as written.
4. Ignore the "Signature" column completely. 
5. Standardize any present times to 12-hour format (e.g., "7:15 AM", "4:30 PM").
6. Capture the "Remarks" exactly as written on that line. If blank, output "".

Output ONLY a valid JSON object formatted exactly like this example:
{
  "rows": [
    {
      "Date": "February 2, 2026",
      "Time_In": "",
      "Time_Out": "",
      "Remarks": "ONLINE MTG HOLIDAY"
    }
  ]
}
"""

client = genai.Client(api_key=API_KEY)

print("Connecting to Gemini API...")

try:
    response = client.models.generate_content(
        model='gemini-2.5-flash', 
        contents=[
            types.Part.from_bytes(data=image_bytes, mime_type='image/jpeg'),
            instructions
        ],
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            temperature=0.1
        )
    )
    
    full_response_string = response.text
    print("Gemini successfully scanned the document!")
    
except Exception as e:
    print(f"\nAPI Error: {e}")
    print("Exiting so you can retry manually.")
    sys.exit(1)


print("Generating Excel File...")
try:
    response_json = json.loads(full_response_string)
    table_rows = response_json.get("rows", [])

    if not isinstance(table_rows, list) or len(table_rows) == 0:
        raise ValueError("AI did not return a valid list of rows.")

    row_dict = {}
    valid_dates = []
    
    for row in table_rows:
        date_str = row.get("Date", "").strip()
        dt = None
        for fmt in ("%B %d, %Y", "%B %b %d, %Y", "%b %d, %Y", "%B %d"):
            try:
                # Default to current or relevant year if missing
                if "," not in date_str and "2026" not in date_str:
                    dt = datetime.strptime(f"{date_str}, 2026", "%B %d, %Y")
                else:
                    dt = datetime.strptime(date_str, fmt)
                break
            except ValueError:
                continue
                
        if dt:
            row_dict[dt.date()] = row
            valid_dates.append(dt.date())
            
    continuous_rows = []
    start_date_str = ""
    end_date_str = ""
    
    if valid_dates:
        start_date = min(valid_dates)
        end_date = max(valid_dates)
        current_date = start_date
        
        start_date_str = start_date.strftime("%B %d")
        end_date_str = end_date.strftime("%B %d, %Y")
        
        while current_date <= end_date:
            if current_date in row_dict:
                continuous_rows.append((current_date, row_dict[current_date]))
            else:
                blank_row = {
                    "Date": current_date.strftime("%B %d, %Y"),
                    "Time_In": "",
                    "Time_Out": "",
                    "Remarks": ""
                }
                continuous_rows.append((current_date, blank_row))
            current_date += timedelta(days=1)
    else:
        continuous_rows = [(None, r) for r in table_rows]
        start_date_str = "Unknown"
        end_date_str = "Unknown"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DTR"

    regular_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    header_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid") 
    weekend_fill = PatternFill(start_color="C6D9F1", end_color="C6D9F1", fill_type="solid")
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                         top=Side(style='thin'), bottom=Side(style='thin'))
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    ws['A1'] = "Month:"
    ws['C1'] = f"{start_date_str} - {end_date_str}" 
    ws['A2'] = "Name:"
    ## Edit this line as we move forward to the project
    ws['C2'] = "May Angela Razon"
    ws['A3'] = "Designation:"
    ws['C3'] = "Project Admin Officer I"
    
    for row in range(1, 4):
        ws[f'A{row}'].font = regular_font
        ws[f'C{row}'].font = regular_font
    ws.merge_cells('C1:E1')
    ws.merge_cells('C2:E2')
    ws.merge_cells('C3:E3')

    headers = ["Date", "Time In", "Time Out", "Absent", "Late", "Remarks"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num)
        cell.value = header
        cell.font = bold_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    start_row = 6
    for i, (dt, row) in enumerate(continuous_rows):
        current_row = start_row + i
        
        date_val = row.get("Date", "")
        t_in = row.get("Time_In", "").strip()
        t_out = row.get("Time_Out", "").strip()
        remarks = row.get("Remarks", "").strip()

        day_name = dt.strftime("%A").upper() if dt else ""

        if not t_in and not t_out and day_name in ["SATURDAY", "SUNDAY"]:
            remarks = day_name

        is_weekend_or_holiday = (not t_in and not t_out and remarks.upper() in ["SUNDAY", "SATURDAY", "HOLIDAY", "ONLINE MTG HOLIDAY"])

        if not remarks and not is_weekend_or_holiday:
            remarks = "Regular Working Days"

        if is_weekend_or_holiday:
            cell_date = ws.cell(row=current_row, column=1)
            cell_date.value = date_val
            cell_date.alignment = right_align
            
            ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=6)
            cell_remark = ws.cell(row=current_row, column=2)
            cell_remark.value = remarks.upper()
            cell_remark.alignment = center_align
            
            for col in range(1, 7):
                c = ws.cell(row=current_row, column=col)
                c.fill = weekend_fill
                c.border = thin_border
                c.font = regular_font
        else:
            row_data = [date_val, t_in, t_out, "0:00", "0:00", remarks]
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_num)
                cell.value = value
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = center_align if col_num > 1 else right_align

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 22

    excel_filename = "DTR_Export.xlsx"
    wb.save(excel_filename)
    print(f"Exported to '{excel_filename}'.")

except Exception as e:
    print(f"\nError parsing or exporting data: {e}")